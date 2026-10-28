import csv
import io
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from assistant import documents, store, tables


class DocumentTableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.patches = [patch.object(store, "DB_PATH", self.root / "test.sqlite3"), patch.object(documents, "DATA_DIR", self.root), patch.object(tables, "DATA_DIR", self.root)]
        for item in self.patches:
            item.start()
        store.init()
        documents.ensure_schema()
        tables.ensure_schema()
        app = FastAPI()
        app.include_router(documents.router)
        app.include_router(tables.router)
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def test_document_retrieval_and_delete_remove_original_and_index(self):
        response = self.client.post("/api/documents", files={"file": ("../映像の企画.txt", "MV は青い照明を使う。\n納期は 12 月です。".encode("utf-8"))})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertEqual(item["name"], "映像の企画.txt")
        result = documents.search_documents("照明の色", [item["id"]])
        self.assertTrue(result)
        self.assertEqual(result[0]["name"], item["name"])
        self.assertIn("行 1", result[0]["location"])
        self.assertEqual(documents.search_documents("照明", []), [])
        self.assertEqual(self.client.request("DELETE", "/api/documents/" + item["id"], json={"confirm": False}).status_code, 400)
        self.assertEqual(self.client.request("DELETE", "/api/documents/" + item["id"], json={"confirm": True}).status_code, 200)
        self.assertFalse((self.root / "documents" / item["id"]).exists())
        self.assertEqual(documents.search_documents("照明"), [])

    def test_bad_format_and_image_pdf_rejected(self):
        self.assertEqual(self.client.post("/api/documents", files={"file": ("legacy.xls", b"abc")}).status_code, 400)
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        data = io.BytesIO()
        writer.write(data)
        response = self.client.post("/api/documents", files={"file": ("scan.pdf", data.getvalue())})
        self.assertEqual(response.status_code, 400)
        self.assertIn("OCR", response.json()["detail"])

    def test_malformed_pdf_has_japanese_error_and_leaves_no_data(self):
        response = self.client.post("/api/documents", files={"file": ("broken.pdf", b"%PDF-1.7\nnot a real document")})
        self.assertEqual(response.status_code, 400)
        self.assertIn("資料を読み込めません", response.json()["detail"])
        self.assertEqual(documents.list_documents(), [])
        self.assertEqual(list((self.root / "documents").iterdir()), [])

    def test_excel_sheet_citation(self):
        workbook = Workbook()
        workbook.active.title = "撮影予定"
        workbook.active.append(["素材", "日数"])
        workbook.active.append(["街の風景", 3])
        data = io.BytesIO()
        workbook.save(data)
        response = self.client.post("/api/documents", files={"file": ("予定.xlsx", data.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        result = documents.search_documents("風景", [response.json()["id"]])
        self.assertIn("撮影予定", result[0]["location"])

    def test_exports_formula_injection_bom_chart_and_original_values(self):
        spec = {"title": "機材比較", "columns": ["機材", "価格", "メモ"], "rows": [["カメラ", 1200, "=HYPERLINK(\"https://example.com\")"], ["マイク", 500, " \t+SUM(1,2)"]], "chart": {"type": "bar", "x": 0, "y": 1, "title": "価格比較", "unit": "円"}}
        response = self.client.post("/api/tables", json=spec)
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertEqual(len(item["files"]), 3)
        path = self.root / "exports" / (item["id"] + ".xlsx")
        wb = load_workbook(path, data_only=False)
        self.assertEqual(wb.active["B2"].value, 1200)
        self.assertTrue(wb.active["C2"].value.startswith("'="))
        self.assertNotEqual(wb.active["C2"].data_type, "f")
        self.assertTrue(wb.active["C3"].value.startswith("'"))
        self.assertEqual(len(wb.active._charts), 1)
        wb.close()
        downloaded = self.client.get(f"/api/files/{item['id']}/csv")
        self.assertEqual(downloaded.status_code, 200)
        self.assertTrue(downloaded.content.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(downloaded.content.decode("utf-8-sig"))))
        self.assertTrue(rows[1][2].startswith("'="))
        svg = self.client.get(f"/api/files/{item['id']}/svg").text
        self.assertIn("価格比較", svg)
        self.assertIn("カメラ: 1200", svg)
        self.assertEqual(self.client.get("/api/files").json()[0]["rows"], spec["rows"])

    def test_graph_rejects_unknown_or_negative_pie(self):
        for kind, values in (("bar", ["不明", 2]), ("pie", [-1, 3])):
            response = self.client.post("/api/tables", json={"columns": ["月", "値"], "rows": [["1月", values[0]], ["2月", values[1]]], "chart": {"type": kind}})
            self.assertEqual(response.status_code, 422)

    def test_import_aggregate_and_sort_without_modifying_source(self):
        original = "月,売上\n1月,100\n2月,300\n1月,200\n2月,不明\n".encode("utf-8-sig")
        response = self.client.post("/api/tables/import", files={"file": ("売上.csv", original)})
        self.assertEqual(response.status_code, 200, response.text)
        preview = response.json()
        preview.update(group_by="月", aggregations=[{"column": "売上", "operation": "sum"}, {"column": "売上", "operation": "average"}], sort_by="月", descending=True)
        response = self.client.post("/api/tables/transform", json=preview)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result["rows"], [["2月", 300.0, 300.0], ["1月", 300.0, 150.0]])
        self.assertIn("除外", result["notes"][-1])
        self.assertEqual(self.client.get("/api/files").json(), [])
        self.assertIn("不明", original.decode("utf-8-sig"))

    def test_exports_persist_and_delete_requires_confirmation(self):
        item = tables.create_table({"columns": ["品名", "数量"], "rows": [["照明", 2]]})
        self.assertEqual(tables.list_files()[0]["id"], item["id"])
        self.assertEqual(self.client.request("DELETE", "/api/files/" + item["id"], json={}).status_code, 400)
        self.assertEqual(self.client.request("DELETE", "/api/files/" + item["id"], json={"confirm": True}).status_code, 200)
        self.assertEqual(tables.list_files(), [])
        self.assertFalse((self.root / "exports" / (item["id"] + ".xlsx")).exists())

    def test_monthly_aggregation_accepts_unambiguous_dates(self):
        response = self.client.post("/api/tables/transform", json={
            "title": "月別売上", "columns": ["日付", "売上"],
            "rows": [["2026-01-02", 100], ["2026/1/31", 200], ["2026年2月1日", 400], ["2026-02-02T13:30:00+09:00", 50]],
            "group_by": "日付", "period": "month", "aggregations": [{"column": "売上", "operation": "sum"}], "sort_by": "日付"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["rows"], [["2026-01", 300.0], ["2026-02", 450.0]])
        self.assertIn("月別", response.json()["notes"][-1])

    def test_monthly_aggregation_rejects_missing_year_and_invalid_date(self):
        for value in ("1/2", "2026/2/30", "2026-13-01", "不明", None, 46024):
            response = self.client.post("/api/tables/transform", json={
                "columns": ["日付", "売上"], "rows": [[value, 100]],
                "group_by": "日付", "period": "month", "aggregations": [{"column": "売上", "operation": "sum"}]})
            self.assertEqual(response.status_code, 400, (value, response.text))
            self.assertIn("2 行目", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
