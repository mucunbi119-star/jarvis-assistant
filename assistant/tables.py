"""User-reviewable tables, safe spreadsheet exports and editable charts."""
from __future__ import annotations

import csv
import io
import json
import math
import numbers
import re
import uuid
from datetime import date, datetime, time, timezone
from html import escape

from fastapi import APIRouter, Body, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, model_validator

from .config import DATA_DIR
from .store import connect
from .documents import MAX_FILE, check_xlsx, decode_text

router = APIRouter()


class ChartSpec(BaseModel):
    type: str = "bar"
    x: int = 0
    y: int = 1
    title: str = "グラフ"
    unit: str = ""


class TableSpec(BaseModel):
    title: str = Field(default="作成した表", min_length=1, max_length=120)
    columns: list[str] = Field(min_length=1, max_length=100)
    rows: list[list[str | int | float | bool | None]] = Field(max_length=10000)
    chart: ChartSpec | None = None
    notes: list[str] = Field(default_factory=list)

    @model_validator(mode="after")
    def validate_shape(self):
        if any(not c or len(c) > 200 for c in self.columns) or len(set(self.columns)) != len(self.columns):
            raise ValueError("列名は空欄にせず、重複のない 200 文字以内の名前を付けてください。")
        if any(len(row) != len(self.columns) for row in self.rows):
            raise ValueError("すべての行の列数を見出しとそろえてください。")
        for row in self.rows:
            for cell in row:
                if isinstance(cell, float) and not math.isfinite(cell):
                    raise ValueError("無限大や NaN は利用できません。")
                if isinstance(cell, str) and len(cell) > 32000:
                    raise ValueError("1 セルは 32,000 文字までです。")
        if self.chart:
            c = self.chart
            if c.type not in ("bar", "line", "pie") or c.x == c.y or not (0 <= c.x < len(self.columns) and 0 <= c.y < len(self.columns)):
                raise ValueError("棒・折れ線・円グラフを選び、異なる有効な X 列・Y 列を指定してください。")
            if not self.rows or len(self.rows) > 100:
                raise ValueError("グラフは 1〜100 行の表から作成してください。")
            values = [_number(row[c.y]) for row in self.rows]
            if any(value is None for value in values):
                raise ValueError("グラフの Y 列はすべて数値が必要です。不明な値を補わず、データを確認してください。")
            if c.type == "pie" and (any(value < 0 for value in values) or sum(values) <= 0):
                raise ValueError("円グラフには負数を使用できず、合計が 0 より大きい必要があります。")
        return self


def ensure_schema():
    (DATA_DIR / "exports").mkdir(parents=True, exist_ok=True)
    with connect() as db:
        db.execute("CREATE TABLE IF NOT EXISTS jarvis_exports (id TEXT PRIMARY KEY, title TEXT NOT NULL, payload_json TEXT NOT NULL, created_at TEXT NOT NULL)")


def _number(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, numbers.Real):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, str):
        try:
            number = float(value.strip())
            return number if math.isfinite(number) else None
        except ValueError:
            pass
    return None


def _safe_cell(value):
    if value is None or value == "":
        return "不明"
    if isinstance(value, str):
        # Excel/CSV formula injection includes leading spaces and control chars.
        clean = "".join(ch for ch in value if ord(ch) >= 32 or ch in "\t\n\r")
        if clean.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) or clean.startswith(("\t", "\r")):
            return "'" + clean
        return clean
    return value


def _svg(spec: TableSpec) -> str:
    c = spec.chart
    labels = [str(row[c.x])[:24] for row in spec.rows]
    values = [_number(row[c.y]) for row in spec.rows]
    width, height = 900, 520
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#07121f"/><g fill="#d9e8ff" font-family="sans-serif"><text x="40" y="38" font-size="22">{escape(c.title)}</text><text x="40" y="65" font-size="13">単位: {escape(c.unit or "指定なし")} ／ {escape(spec.columns[c.y])}</text>']
    colors = ["#38bdf8", "#818cf8", "#2dd4bf", "#fbbf24", "#fb7185", "#c084fc"]
    if c.type == "pie":
        total, angle = sum(values), -math.pi / 2
        for index, (label, value) in enumerate(zip(labels, values)):
            end = angle + value / total * 2 * math.pi
            x1, y1 = 270 + math.cos(angle) * 160, 285 + math.sin(angle) * 160
            x2, y2 = 270 + math.cos(end) * 160, 285 + math.sin(end) * 160
            color = colors[index % len(colors)]
            if value == total:
                parts.append(f'<circle cx="270" cy="285" r="160" fill="{color}"/>')
            elif value:
                parts.append(f'<path d="M270 285 L{x1:.3f} {y1:.3f} A160 160 0 {int(end-angle > math.pi)} 1 {x2:.3f} {y2:.3f} Z" fill="{color}"/>')
            # Scroll-free small legend; for large charts, detailed table is primary.
            if index < 18:
                parts.append(f'<rect x="490" y="{100 + index * 20}" width="12" height="12" fill="{color}"/><text x="510" y="{111 + index * 20}" font-size="12">{escape(label)}: {value:g} ({value/total*100:.1f}%)</text>')
            angle = end
    else:
        low, high = min(0, min(values)), max(0, max(values))
        if low == high:
            high = low + 1
        def y(value):
            return 430 - (value - low) / (high - low) * 320
        for tick in range(6):
            value = low + (high - low) * tick / 5
            yy = y(value)
            parts.append(f'<path d="M80 {yy:.3f} H850" stroke="#26384d"/><text x="70" y="{yy + 4:.3f}" text-anchor="end" font-size="12">{value:g}</text>')
        parts.append(f'<path d="M80 {y(0):.3f} H850" stroke="#a7bed8"/>')
        step = 770 / len(values)
        points = []
        for index, (label, value) in enumerate(zip(labels, values)):
            x = 80 + (index + .5) * step
            yy = y(value)
            if c.type == "bar":
                parts.append(f'<rect x="{x-step*.35:.3f}" y="{min(yy,y(0)):.3f}" width="{step*.7:.3f}" height="{abs(yy-y(0)):.3f}" fill="#38bdf8"><title>{escape(label)}: {value:g}</title></rect>')
            else:
                points.append(f"{x:.3f},{yy:.3f}")
                parts.append(f'<circle cx="{x:.3f}" cy="{yy:.3f}" r="4" fill="#38bdf8"><title>{escape(label)}: {value:g}</title></circle>')
            if index % max(1, len(values)//10) == 0:
                parts.append(f'<text x="{x:.3f}" y="455" text-anchor="middle" font-size="11">{escape(label[:12])}</text>')
        if points:
            parts.append(f'<polyline points="{" ".join(points)}" fill="none" stroke="#38bdf8" stroke-width="2"/>')
        parts.append(f'<text x="470" y="490" text-anchor="middle" font-size="13">{escape(spec.columns[c.x])}</text>')
    parts.append("</g></svg>")
    return "".join(parts)


def create_table(spec: TableSpec | dict):
    if isinstance(spec, dict):
        spec = TableSpec.model_validate(spec)
    ensure_schema()
    identifier = uuid.uuid4().hex
    export_dir = DATA_DIR / "exports"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "データ"
    sheet.append([_safe_cell(value) for value in spec.columns])
    for row in spec.rows:
        sheet.append([_safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="123D69")
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in list(column)[:200])
        sheet.column_dimensions[column[0].column_letter].width = min(45, max(14, width * 1.4))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.########'
    if spec.chart:
        c = spec.chart
        # Only validated numeric values are inserted into the chart series.
        for index, row in enumerate(spec.rows, 2):
            sheet.cell(index, c.y + 1, _number(row[c.y]))
        chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[c.type]()
        chart.title = c.title
        if c.type != "pie":
            chart.y_axis.title = c.unit or spec.columns[c.y]
            chart.x_axis.title = spec.columns[c.x]
            chart.y_axis.scaling.min = min(0, min(_number(row[c.y]) for row in spec.rows))
        chart.add_data(Reference(sheet, min_col=c.y + 1, min_row=1, max_row=len(spec.rows) + 1), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=c.x + 1, min_row=2, max_row=len(spec.rows) + 1))
        sheet.add_chart(chart, "A" + str(len(spec.rows) + 4))
    if spec.notes:
        notes = workbook.create_sheet("注記")
        notes.column_dimensions["A"].width = 100
        for note in spec.notes:
            notes.append([_safe_cell(note)])
    created = datetime.now(timezone.utc).isoformat()
    payload = spec.model_dump()
    payload.update({"id": identifier, "created_at": created})
    formats = ["xlsx", "csv"] + (["svg"] if spec.chart else [])
    payload["files"] = [{"name": spec.title + "." + kind, "url": f"/api/files/{identifier}/{kind}", "format": kind} for kind in formats]
    try:
        workbook.save(export_dir / (identifier + ".xlsx"))
        with (export_dir / (identifier + ".csv")).open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow([_safe_cell(value) for value in spec.columns])
            writer.writerows([[_safe_cell(value) for value in row] for row in spec.rows])
        if spec.chart:
            (export_dir / (identifier + ".svg")).write_text(_svg(spec), encoding="utf-8")
        with connect() as db:
            db.execute("INSERT INTO jarvis_exports VALUES (?,?,?,?)", (identifier, spec.title, json.dumps(payload, ensure_ascii=False), created))
    except Exception:
        for kind in formats:
            (export_dir / (identifier + "." + kind)).unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return payload


@router.post("/api/tables")
def tables_create(spec: TableSpec):
    return create_table(spec)


@router.get("/api/files")
def list_files():
    ensure_schema()
    with connect() as db:
        return [json.loads(row["payload_json"]) for row in db.execute("SELECT payload_json FROM jarvis_exports ORDER BY created_at DESC")]


@router.get("/api/files/{identifier}/{kind}")
def download_file(identifier: str, kind: str):
    if kind not in ("csv", "xlsx", "svg"):
        raise HTTPException(404, "ファイル形式が見つかりません。")
    with connect() as db:
        row = db.execute("SELECT * FROM jarvis_exports WHERE id=?", (identifier,)).fetchone()
    if not row:
        raise HTTPException(404, "作成ファイルが見つかりません。")
    path = DATA_DIR / "exports" / (row["id"] + "." + kind)
    if not path.is_file():
        raise HTTPException(404, "作成ファイルが見つかりません。")
    media = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8", "svg": "image/svg+xml"}[kind]
    return FileResponse(path, media_type=media, filename=row["title"] + "." + kind)


@router.delete("/api/files/{identifier}")
def delete_file(identifier: str, body: dict = Body(...)):
    if body.get("confirm") is not True:
        raise HTTPException(400, "作成した表と出力ファイルの削除を確認してください。")
    with connect() as db:
        row = db.execute("SELECT id FROM jarvis_exports WHERE id=?", (identifier,)).fetchone()
        if not row:
            raise HTTPException(404, "作成ファイルが見つかりません。")
        for kind in ("xlsx", "csv", "svg"):
            (DATA_DIR / "exports" / (row["id"] + "." + kind)).unlink(missing_ok=True)
        db.execute("DELETE FROM jarvis_exports WHERE id=?", (identifier,))
    return {"deleted": True}


@router.post("/api/tables/import")
async def import_table(file: UploadFile = File(...)):
    data = await file.read(MAX_FILE + 1)
    await file.close()
    if len(data) > MAX_FILE:
        raise HTTPException(413, "ファイルは 10 MB までです。")
    name = file.filename or ""
    try:
        if name.lower().endswith(".csv"):
            rows = list(csv.reader(io.StringIO(decode_text(data))))
        elif name.lower().endswith(".xlsx"):
            check_xlsx(data)
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                sheet = workbook.active
                if sheet.max_row > 10001 or sheet.max_column > 100:
                    raise ValueError("10,000 行・100 列まで読み込めます。")
                rows = [[v.isoformat() if isinstance(v, (date, datetime, time)) else v for v in row] for row in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
        else:
            raise ValueError("表の読み込みは CSV・Excel（.xlsx）に対応しています。")
        if len(rows) < 2:
            raise ValueError("見出しと 1 行以上のデータが必要です。")
        width = max(len(row) for row in rows)
        columns = [str(value) if value not in (None, "") else f"列 {index + 1}" for index, value in enumerate(rows[0] + [None] * (width - len(rows[0])))]
        if len(set(columns)) != len(columns):
            columns = [f"{value} ({index + 1})" for index, value in enumerate(columns)]
        spec = TableSpec(title=name.rsplit(".", 1)[0][:120] or "読み込んだ表", columns=columns, rows=[row + [None] * (width - len(row)) for row in rows[1:]], notes=["元ファイルは変更していません。Excel の数式は保存済み計算値を読み込み、未計算値は不明になります。"])
        return spec.model_dump()
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from None
    except Exception:
        raise HTTPException(400, "表を読み込めませんでした。形式・ファイルの破損を確認してください。") from None


class Aggregation(BaseModel):
    column: str
    operation: str = "sum"


class TransformSpec(TableSpec):
    group_by: str | None = None
    period: str | None = None
    aggregations: list[Aggregation] = Field(default_factory=list, max_length=100)
    sort_by: str | None = None
    descending: bool = False


def _month(value, row_number):
    """Only year-first, unambiguous calendar dates are accepted."""
    if not isinstance(value, str):
        raise HTTPException(400, f"{row_number} 行目の日付を月別に変換できません。年を含む YYYY-MM-DD、YYYY/MM/DD、YYYY年M月D日形式にしてください。")
    value = value.strip()
    parsed = None
    try:
        if re.fullmatch(r"\d{4}-\d{2}", value):
            parsed = datetime.strptime(value, "%Y-%m")
        elif re.fullmatch(r"\d{4}/\d{1,2}/\d{1,2}", value):
            parsed = datetime.strptime(value, "%Y/%m/%d")
        elif re.fullmatch(r"\d{4}年\d{1,2}月\d{1,2}日", value):
            parsed = datetime.strptime(value, "%Y年%m月%d日")
        elif re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ].+)?", value):
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        pass
    if parsed is None:
        raise HTTPException(400, f"{row_number} 行目の日付を月別に変換できません。日付が正しいか確認し、年を含む YYYY-MM-DD、YYYY/MM/DD、YYYY年M月D日形式にしてください。年や月は推測しません。")
    return f"{parsed.year:04d}-{parsed.month:02d}"


@router.post("/api/tables/transform")
def transform_table(spec: TransformSpec):
    columns, rows, notes = list(spec.columns), list(spec.rows), list(spec.notes)
    if spec.period not in (None, "month"):
        raise HTTPException(400, "日付の変換は月別（month）に対応しています。")
    if spec.period and not spec.group_by:
        raise HTTPException(400, "月別に集計する日付列を指定してください。")
    if spec.group_by:
        if spec.group_by not in columns or not spec.aggregations:
            raise HTTPException(400, "集計キーと、1 件以上の集計方法を選んでください。")
        for agg in spec.aggregations:
            if agg.column not in columns or agg.operation not in ("sum", "average", "count"):
                raise HTTPException(400, "集計列・集計方法が無効です。合計・平均・件数に対応しています。")
        key_index = columns.index(spec.group_by)
        groups = {}
        for row_number, row in enumerate(rows, 2):
            key = (_month(row[key_index], row_number) if spec.period == "month"
                   else str(row[key_index]) if row[key_index] not in (None, "") else "不明")
            groups.setdefault(key, []).append(row)
        result = []
        ignored = 0
        for key, group in groups.items():
            entry = [key]
            for agg in spec.aggregations:
                index = columns.index(agg.column)
                if agg.operation == "count":
                    entry.append(sum(row[index] not in (None, "", "不明") for row in group))
                else:
                    values = [_number(row[index]) for row in group]
                    ignored += sum(value is None for value in values)
                    values = [value for value in values if value is not None]
                    entry.append((sum(values) if agg.operation == "sum" else sum(values) / len(values)) if values else "不明")
            result.append(entry)
        columns = [spec.group_by] + [f"{agg.column}（{ {'sum':'合計','average':'平均','count':'件数'}[agg.operation]}）" for agg in spec.aggregations]
        rows = result
        if spec.period == "month":
            notes.append(f"「{spec.group_by}」列の日付を YYYY-MM 形式に変換し、月別に集計しました。時刻付き日付は記載された現地日付の月を使用しています。")
        if ignored:
            notes.append(f"集計時、数値として解釈できない {ignored} セルは合計・平均から除外しました。値がない集計結果は「不明」です。")
    if spec.sort_by:
        if spec.sort_by not in columns:
            raise HTTPException(400, "並べ替え列が見つかりません。集計後の列名を指定してください。")
        index = columns.index(spec.sort_by)
        def key(row):
            value = row[index]
            numeric = _number(value)
            return (0, numeric) if numeric is not None else (1, str(value))
        rows.sort(key=key, reverse=spec.descending)
    # Return a preview; saving is a separate, explicit user action.
    return TableSpec(title=spec.title, columns=columns, rows=rows, notes=notes).model_dump()
