"""Local document extraction and bounded, citation-preserving retrieval."""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Body, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import DATA_DIR
from .store import connect

router = APIRouter()
MAX_FILE = 10 * 1024 * 1024
MAX_TEXT = 1_000_000


def ensure_schema():
    (DATA_DIR / "documents").mkdir(parents=True, exist_ok=True)
    with connect() as db:
        db.execute("""CREATE TABLE IF NOT EXISTS jarvis_documents (
            id TEXT PRIMARY KEY, name TEXT NOT NULL, kind TEXT NOT NULL,
            created_at TEXT NOT NULL, size INTEGER NOT NULL, sections_json TEXT NOT NULL)""")


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8", "cp932"):
        try:
            value = data.decode(encoding)
            if "\x00" in value:
                raise ValueError("テキストとして読み込めないファイルです。")
            return value
        except UnicodeDecodeError:
            continue
    raise ValueError("文字コードを判別できません。UTF-8 または日本語 Windows 形式で保存してください。")


def check_xlsx(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("展開後の Excel ファイルが大きすぎます（100 MB まで）。")
    except zipfile.BadZipFile as exc:
        raise ValueError("有効な .xlsx ファイルではありません。") from exc


def extract_sections(name: str, data: bytes) -> tuple[str, list[dict]]:
    kind = Path(name).suffix.lower().lstrip(".")
    sections = []
    if kind == "txt":
        text = decode_text(data)
        # Chunking preserves locations, avoids sending entire documents to an API.
        lines = text.splitlines()
        for index in range(0, len(lines), 40):
            sections.append({"location": f"行 {index + 1}〜{min(index + 40, len(lines))}", "text": "\n".join(lines[index:index + 40])})
    elif kind == "csv":
        rows = list(csv.reader(io.StringIO(decode_text(data))))
        if len(rows) > 10000 or any(len(row) > 100 for row in rows):
            raise ValueError("CSV は 10,000 行・100 列まで読み込めます。")
        for index in range(0, len(rows), 30):
            sections.append({"location": f"行 {index + 1}〜{min(index + 30, len(rows))}", "text": "\n".join(" | ".join(row) for row in rows[index:index + 30])})
    elif kind == "pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("暗号化 PDF は未対応です。解除した別ファイルをご用意ください。")
        if len(reader.pages) > 300:
            raise ValueError("PDF は 300 ページまで読み込めます。")
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                sections.append({"location": f"ページ {index + 1}", "text": text})
        if not sections:
            raise ValueError("抽出できる文字がありません。画像だけの PDF（OCR）は未対応です。")
    elif kind == "xlsx":
        check_xlsx(data)
        # Never execute formulas, macros or external links.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                if sheet.max_row > 10000 or sheet.max_column > 100:
                    raise ValueError("各シートは 10,000 行・100 列まで読み込めます。")
                rows = list(sheet.iter_rows(values_only=True))
                for index in range(0, len(rows), 30):
                    sections.append({"location": f"シート「{sheet.title}」行 {index + 1}〜{min(index + 30, len(rows))}", "text": "\n".join(" | ".join("不明" if value is None else str(value) for value in row) for row in rows[index:index + 30])})
        finally:
            workbook.close()
    else:
        raise ValueError("対応形式は TXT・CSV・文字抽出可能な PDF・Excel（.xlsx）です。.xls・画像・OCR は未対応です。")
    sections = [section for section in sections if section["text"].strip()]
    if not sections:
        raise ValueError("読み込めるテキストがありません。")
    if sum(len(s["text"]) for s in sections) > MAX_TEXT:
        raise ValueError("抽出テキストは合計 100 万文字までです。小さな資料に分けてください。")
    return kind, sections


def _public(row):
    value = dict(row)
    value["sections"] = len(json.loads(value.pop("sections_json")))
    return value


def list_documents():
    ensure_schema()
    with connect() as db:
        return [_public(row) for row in db.execute("SELECT * FROM jarvis_documents ORDER BY created_at DESC")]


@router.get("/api/documents")
def documents_list():
    return list_documents()


@router.post("/api/documents")
async def upload_document(file: UploadFile = File(...)):
    ensure_schema()
    data = await file.read(MAX_FILE + 1)
    await file.close()
    if len(data) > MAX_FILE:
        raise HTTPException(413, "資料は 1 ファイル 10 MB までです。")
    name = (file.filename or "資料.txt").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f]", "", name)[:180]
    try:
        kind, sections = extract_sections(name, data)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from None
    except Exception:
        raise HTTPException(400, "資料を読み込めませんでした。ファイルが破損していないか、形式を確認してください。") from None
    identifier = uuid.uuid4().hex
    created = datetime.now(timezone.utc).isoformat()
    path = DATA_DIR / "documents" / identifier
    try:
        path.write_bytes(data)
        with connect() as db:
            db.execute("INSERT INTO jarvis_documents VALUES (?,?,?,?,?,?)", (identifier, name, kind, created, len(data), json.dumps(sections, ensure_ascii=False)))
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return {"id": identifier, "name": name, "kind": kind, "created_at": created, "size": len(data), "sections": len(sections), "note": "ローカルに保存しました。実サービスで質問すると、関連する抜粋のみ OpenAI API に送信します。"}


@router.delete("/api/documents/{identifier}")
def delete_document(identifier: str, body: dict = Body(...)):
    if body.get("confirm") is not True:
        raise HTTPException(400, "資料と検索データの削除を確認してください。")
    with connect() as db:
        row = db.execute("SELECT id FROM jarvis_documents WHERE id=?", (identifier,)).fetchone()
        if row is None:
            raise HTTPException(404, "資料が見つかりません。")
        # Use the validated stored id, never an arbitrary incoming path.
        (DATA_DIR / "documents" / row["id"]).unlink(missing_ok=True)
        db.execute("DELETE FROM jarvis_documents WHERE id=?", (identifier,))
    return {"deleted": True}


def _terms(value: str) -> set[str]:
    value = unicodedata.normalize("NFKC", value).casefold()
    words = set(re.findall(r"[a-z0-9]{2,}", value))
    for run in re.findall(r"[ぁ-んァ-ヶ一-龠]{2,}", value):
        words.update(run[index:index + 2] for index in range(len(run) - 1))
    return words


def search_documents(query: str, ids: list[str] | None = None) -> list[dict]:
    """Retrieve at most 6 small, relevant passages; never invent a reference."""
    ensure_schema()
    terms = _terms(query)
    if not terms:
        return []
    with connect() as db:
        rows = db.execute("SELECT * FROM jarvis_documents").fetchall()
    ranked = []
    for row in rows:
        if ids is not None and row["id"] not in ids:
            continue
        for section in json.loads(row["sections_json"]):
            # Rank paragraphs within a long page so citations remain useful.
            text = section["text"]
            for start in range(0, len(text), 1200):
                excerpt = text[start:start + 1500]
                score = len(terms & _terms(excerpt + " " + row["name"]))
                if score:
                    ranked.append((score, {"id": row["id"], "name": row["name"], "title": row["name"], "location": section["location"], "text": excerpt}))
    ranked.sort(key=lambda item: item[0], reverse=True)
    return [item[1] for item in ranked[:6]]
